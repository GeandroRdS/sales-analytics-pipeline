from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class ExcelReportGenerator:
    """
    Generates the operational Excel report containing
    processing summary, approved orders and rejection details.
    """

    def generate(
        self,
        results: list[dict],
        parsing_errors: list[dict],
        output_path: Path,
    ) -> None:

        output_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        approved_rows = []
        rejected_order_rows = []
        rejected_item_rows = []

        total_items = 0

        rejection_counter = {}

        for result in results:

            total_items += len(
                result["items"]
            )

            if result["status"] == "APPROVED":

                self._append_approved_rows(
                    approved_rows,
                    result,
                )

            else:

                reasons = ", ".join(
                    result["errors"]
                )

                rejected_order_rows.append(
                    {
                        "purchase_order": result[
                            "purchase_order"
                        ],
                        "order_date": result[
                            "order_date"
                        ],
                        "customer_code": result[
                            "customer_code"
                        ],
                        "customer_name": result[
                            "customer_name"
                        ],
                        "status": result[
                            "status"
                        ],
                        "rejection_reasons": reasons,
                    }
                )

                for error in result["errors"]:
                    rejection_counter[error] = (
                        rejection_counter.get(
                            error,
                            0,
                        )
                        + 1
                    )

                self._append_rejected_items(
                    rejected_item_rows,
                    result,
                )

        approved_orders = sum(
            1
            for result in results
            if result["status"] == "APPROVED"
        )

        rejected_orders = (
            len(results)
            - approved_orders
        )

        approval_rate = (
            approved_orders
            / len(results)
            if results
            else 0
        )

        summary_rows = [
            {
                "metric": "Orders Processed",
                "value": len(results),
            },
            {
                "metric": "Approved Orders",
                "value": approved_orders,
            },
            {
                "metric": "Rejected Orders",
                "value": rejected_orders,
            },
            {
                "metric": "Approval Rate",
                "value": approval_rate,
            },
            {
                "metric": "Items Processed",
                "value": total_items,
            },
            {
                "metric": "Parsing Errors",
                "value": len(parsing_errors),
            },
        ]

        rejection_rows = [
            {
                "rejection_reason": reason,
                "affected_orders": quantity,
            }
            for reason, quantity
            in sorted(
                rejection_counter.items(),
                key=lambda item: item[1],
                reverse=True,
            )
        ]

        summary_df = pd.DataFrame(
            summary_rows
        )

        rejection_summary_df = pd.DataFrame(
            rejection_rows
        )

        approved_df = pd.DataFrame(
            approved_rows
        )

        rejected_orders_df = pd.DataFrame(
            rejected_order_rows
        )

        rejected_items_df = pd.DataFrame(
            rejected_item_rows
        )

        with pd.ExcelWriter(
            output_path,
            engine="openpyxl",
        ) as writer:

            summary_df.to_excel(
                writer,
                sheet_name="Processing Summary",
                index=False,
                startrow=1,
            )

            rejection_summary_df.to_excel(
                writer,
                sheet_name="Processing Summary",
                index=False,
                startrow=10,
            )

            approved_df.to_excel(
                writer,
                sheet_name="Approved Orders",
                index=False,
            )

            rejected_orders_df.to_excel(
                writer,
                sheet_name="Rejected Orders",
                index=False,
            )

            rejected_items_df.to_excel(
                writer,
                sheet_name="Rejected Items",
                index=False,
            )

        self._format_workbook(
            output_path
        )

    def _append_approved_rows(
        self,
        rows: list[dict],
        result: dict,
    ) -> None:

        customer = (
            result["customer"]
            or {}
        )

        seller = (
            result["seller"]
            or {}
        )

        for item in result["items"]:

            rows.append(
                {
                    "purchase_order": result[
                        "purchase_order"
                    ],
                    "order_date": result[
                        "order_date"
                    ],
                    "customer_code": result[
                        "customer_code"
                    ],
                    "customer_name": result[
                        "customer_name"
                    ],
                    "region": customer.get(
                        "region"
                    ),
                    "seller_code": seller.get(
                        "seller_code"
                    ),
                    "seller_name": seller.get(
                        "seller_name"
                    ),
                    "product_code": item[
                        "product_code"
                    ],
                    "description": item[
                        "description"
                    ],
                    "category": item[
                        "category"
                    ],
                    "brand": item[
                        "brand"
                    ],
                    "quantity": item[
                        "quantity"
                    ],
                    "units_per_case": item[
                        "units_per_case"
                    ],
                    "unit_price": item[
                        "unit_price"
                    ],
                    "line_total": item[
                        "line_total"
                    ],
                }
            )

    def _append_rejected_items(
        self,
        rows: list[dict],
        result: dict,
    ) -> None:

        for item in result["items"]:

            if not item["errors"]:
                continue

            rows.append(
                {
                    "purchase_order": result[
                        "purchase_order"
                    ],
                    "customer_code": result[
                        "customer_code"
                    ],
                    "customer_name": result[
                        "customer_name"
                    ],
                    "product_code": item[
                        "product_code"
                    ],
                    "description": item[
                        "description"
                    ],
                    "quantity": item[
                        "quantity"
                    ],
                    "stock_quantity": item[
                        "stock_quantity"
                    ],
                    "units_per_case": item[
                        "units_per_case"
                    ],
                    "errors": ", ".join(
                        item["errors"]
                    ),
                }
            )

    def _format_workbook(
        self,
        output_path: Path,
    ) -> None:

        workbook = load_workbook(
            output_path
        )

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        header_font = Font(
            color="FFFFFF",
            bold=True,
        )

        for worksheet in workbook.worksheets:

            worksheet.freeze_panes = "A2"

            worksheet.auto_filter.ref = (
                worksheet.dimensions
            )

            for cell in worksheet[1]:

                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(
                    horizontal="center"
                )

            for column_cells in (
                worksheet.columns
            ):

                max_length = 0

                column_letter = (
                    get_column_letter(
                        column_cells[0].column
                    )
                )

                for cell in column_cells:

                    value = (
                        ""
                        if cell.value is None
                        else str(cell.value)
                    )

                    max_length = max(
                        max_length,
                        len(value),
                    )

                worksheet.column_dimensions[
                    column_letter
                ].width = min(
                    max_length + 2,
                    45,
                )

        # Summary formatting
        summary_sheet = workbook[
            "Processing Summary"
        ]

        summary_sheet["A1"] = (
            "ORDER PROCESSING SUMMARY"
        )

        summary_sheet["A1"].font = Font(
            bold=True,
            size=14,
        )

        # Second table header
        for cell in summary_sheet[11]:

            cell.fill = header_fill
            cell.font = header_font

        summary_sheet["B5"].number_format = (
            "0.00%"
        )

        # Currency formatting
        approved_sheet = workbook[
            "Approved Orders"
        ]

        header_map = {
            cell.value: cell.column
            for cell in approved_sheet[1]
        }

        for column_name in (
            "unit_price",
            "line_total",
        ):

            column_index = header_map.get(
                column_name
            )

            if column_index:

                for row in range(
                    2,
                    approved_sheet.max_row + 1,
                ):

                    approved_sheet.cell(
                        row=row,
                        column=column_index,
                    ).number_format = (
                        '#,##0.00'
                    )

        workbook.save(
            output_path
        )